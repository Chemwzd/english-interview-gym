"""EngTraining FastAPI 入口。

启动：cd app && ../.venv/bin/python -m uvicorn server.main:app --host 127.0.0.1 --port 8765
（或直接 bash scripts/run_server.sh）
"""
import re
import shutil
import time
from pathlib import Path

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles

from . import config, prompts, store, tts
from . import session as sess

WEB = Path(__file__).resolve().parents[1] / "web"

app = FastAPI(title="EngTraining", version="0.1.0")


@app.get("/api/health")
def health():
    return {
        "ok": True,
        "llm_model": config.get("llm.model"),
        "asr_driver": config.get("asr.driver"),
        "tts_driver": config.get("tts.driver"),
        "max_answer_seconds": config.get("session.max_answer_seconds", 120),
        "tokenhub_key_set": bool(config.tokenhub_key()),
    }


@app.get("/api/personas")
def personas():
    return prompts.list_personas()


@app.get("/api/question-sets")
def question_sets():
    return prompts.list_question_sets()


@app.post("/api/session/start")
def session_start(payload: dict):
    try:
        return sess.start(payload.get("persona", "hr-friendly"), payload.get("set", "baseline-8"), start_idx=payload.get("start", 0))
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, str(e))


@app.get("/api/session/{sid}")
def session_get(sid: str):
    try:
        return sess.get_state(sid)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(404, str(e))


@app.post("/api/session/{sid}/answer")
async def session_answer(
    sid: str,
    audio: UploadFile = File(...),
    mode: str = Form("free"),
    script: str = Form(""),
):
    d = config.data_dir() / "audio" / sid
    d.mkdir(parents=True, exist_ok=True)
    ext = Path(audio.filename or "a.webm").suffix or ".webm"
    dst = d / f"{int(time.time() * 1000)}{ext}"
    with dst.open("wb") as f:
        shutil.copyfileobj(audio.file, f)
    try:
        return sess.submit_answer(sid, dst, mode=mode, script=script)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, str(e))


@app.get("/api/session/{sid}/suggestion")
def session_suggestion(sid: str, refresh: int = 0):
    try:
        return sess.get_suggestion(sid, refresh=bool(refresh))
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, str(e))


@app.post("/api/session/{sid}/skip")
def session_skip(sid: str):
    try:
        return sess.skip_current(sid)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, str(e))


@app.post("/api/session/{sid}/end")
def session_end(sid: str):
    try:
        return sess.end(sid)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, str(e))


@app.get("/api/sessions")
def sessions_list():
    return store.list_sessions()


@app.get("/api/gloss")
def gloss_get(word: str, context: str = ""):
    try:
        return sess.gloss(word, context)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, str(e))


@app.post("/api/favorite")
def favorite_add(payload: dict):
    try:
        store.append_favorite(
            {
                "type": payload.get("type", "word"),
                "text": (payload.get("text") or "")[:800],
                "note": (payload.get("note") or "")[:400],
                "session": payload.get("session", ""),
            }
        )
        return {"ok": True}
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, str(e))


@app.get("/api/favorites")
def favorites_list():
    return store.list_favorites()


@app.get("/api/stats")
def stats_get():
    return store.stats()


@app.get("/api/tts")
def tts_get(text: str, persona: str = ""):
    try:
        data = tts.synth(text, persona=persona or None)
        return Response(content=data, media_type="audio/mpeg")
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, str(e))


@app.get("/api/set/{set_key}")
def set_detail(set_key: str):
    """题集详情：每题标题/图片/练习统计/历史（供题目卡片墙）。"""
    try:
        qset = prompts.load_question_set(set_key)
        meta = (prompts.load_question_meta() or {}).get(set_key) or {}
        hist = store.attempts_by_question()
        img_dir = config.materials_dir() / "question-bank" / "images" / set_key
        out = []
        for i, q in enumerate(qset.get("questions") or []):
            qid = q.get("id") or f"q{i + 1}"
            qtext = (q.get("text") or "").strip()
            atts = hist.get(qtext, [])
            m = meta.get(qid) or {}
            img = f"/media/{set_key}/{qid}.png" if (img_dir / f"{qid}.png").exists() else None
            best = max([a.get("score") for a in atts if a.get("score")], default=None)
            out.append(
                {
                    "i": i, "id": qid, "text": qtext,
                    "round": q.get("round"), "category": q.get("category"), "tag": q.get("tag"),
                    "intent": q.get("intent"), "hint": q.get("hint"), "title": m.get("title"),
                    "img": img, "attempts": len(atts), "last_at": atts[0].get("ts") if atts else None,
                    "best_score": best, "history": atts[:20],
                }
            )
        done = sum(1 for x in out if x["attempts"] > 0)
        return {"key": set_key, "title": qset.get("title", set_key), "note": qset.get("note", ""),
                "done": done, "count": len(out), "questions": out}
    except Exception as e:  # noqa: BLE001
        raise HTTPException(404, str(e))


def _extract_resume_text(filename: str, raw: bytes) -> str:
    """简历文件 → 纯文本（支持 docx / pdf / txt / md）。"""
    name = (filename or "").lower()
    if name.endswith(".docx"):
        import io

        import docx  # python-docx

        d = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs]
        for tb in d.tables:
            for row in tb.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(x for x in parts if x.strip())
    if name.endswith(".pdf"):
        import io

        import pypdf

        r = pypdf.PdfReader(io.BytesIO(raw))
        return "\n".join((pg.extract_text() or "") for pg in r.pages)
    return raw.decode("utf-8", errors="ignore")


@app.get("/api/profile")
def profile_status():
    """个人资料状态（本地保存；导入后示范答案将额外提供「定制版」）。"""
    t = prompts.load_profile()
    return {"has_profile": bool(t), "chars": len(t), "head": t[:150]}


@app.post("/api/profile/text")
def profile_set_text(payload: dict):
    text = (payload.get("text") or "").strip()
    if len(text) < 30:
        raise HTTPException(400, "内容太短（至少 30 字）")
    (config.materials_dir() / "profile.md").write_text(text, encoding="utf-8")
    return {"ok": True, "chars": len(text)}


@app.post("/api/profile/upload")
async def profile_upload(file: UploadFile = File(...)):
    raw = await file.read()
    if len(raw) > 8 * 1024 * 1024:
        raise HTTPException(400, "文件过大（上限 8MB）")
    try:
        text = _extract_resume_text(file.filename or "", raw).strip()
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"解析失败：{e}")
    if len(text) < 30:
        raise HTTPException(400, "解析出的文本太短，请检查文件或改用粘贴方式")
    (config.materials_dir() / "profile.md").write_text(text, encoding="utf-8")
    return {"ok": True, "chars": len(text)}


@app.delete("/api/profile")
def profile_delete():
    p = config.materials_dir() / "profile.md"
    if p.exists():
        p.unlink()
    return {"ok": True}


@app.get("/")
def index():
    # no-cache + 按资源 mtime 自动注入 ?v= 版本号：改前端后刷新即生效，无需手动改版本
    html = (WEB / "index.html").read_text(encoding="utf-8")
    ver = int(max((WEB / "app.js").stat().st_mtime, (WEB / "style.css").stat().st_mtime))
    html = re.sub(r"\?v=\d+", f"?v={ver}", html)
    return Response(html, media_type="text/html", headers={"Cache-Control": "no-cache"})


app.mount("/static", StaticFiles(directory=str(WEB)), name="static")

# 题目形象图（由 VOD AIGC 生成，本地缓存）
_MEDIA_DIR = config.materials_dir() / "question-bank" / "images"
_MEDIA_DIR.mkdir(parents=True, exist_ok=True)
app.mount("/media", StaticFiles(directory=str(_MEDIA_DIR)), name="media")
